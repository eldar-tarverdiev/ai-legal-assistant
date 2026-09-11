import os
from docx import Document


def create_test_contract():
    os.makedirs("data/contracts", exist_ok=True)

    doc = Document()
    doc.add_heading('ТРУДОВОЙ ДОГОВОР №123', 0)

    doc.add_paragraph(
        'г. Москва                                                                                     01 января 2026 г.')
    doc.add_paragraph('')
    doc.add_paragraph(
        'ООО "ТехноКомпани", именуемое в дальнейшем "Работодатель", в лице генерального директора Иванова И.И., действующего на основании Устава, с одной стороны, и')
    doc.add_paragraph(
        'Петров Петр Петрович, именуемый в дальнейшем "Работник", с другой стороны, заключили настоящий трудовой договор о нижеследующем:')
    doc.add_paragraph('')

    doc.add_heading('1. ПРЕДМЕТ ДОГОВОРА', level=1)
    doc.add_paragraph('1.1. Работник принимается на работу на должность ML-инженера.')
    doc.add_paragraph('1.2. Место работы: г. Москва, офис компании.')
    doc.add_paragraph('')

    doc.add_heading('2. ПРАВА И ОБЯЗАННОСТИ СТОРОН', level=1)
    doc.add_paragraph('2.1. Работник обязан:')
    doc.add_paragraph('- добросовестно исполнять свои трудовые обязанности;')
    doc.add_paragraph('- соблюдать правила внутреннего трудового распорядка;')
    doc.add_paragraph('- разрабатывать ML-модели и внедрять их в продакшен.')
    doc.add_paragraph('')
    doc.add_paragraph('2.2. Работодатель обязан:')
    doc.add_paragraph('- обеспечить Работнику условия работы, предусмотренные трудовым законодательством;')
    doc.add_paragraph('- своевременно и в полном размере выплачивать Работнику заработную плату.')
    doc.add_paragraph('')

    doc.add_heading('3. ОПЛАТА ТРУДА', level=1)
    doc.add_paragraph(
        '3.1. За выполнение трудовых обязанностей Работнику устанавливается должностной оклад в размере 150 000 (сто пятьдесят тысяч) рублей в месяц.')
    doc.add_paragraph('3.2. Заработная плата выплачивается каждые 2 месяца.')
    doc.add_paragraph('')

    doc.add_heading('4. РАБОЧЕЕ ВРЕМЯ И ВРЕМЯ ОТДЫХА', level=1)
    doc.add_paragraph('4.1. Работнику устанавливается ненормированный рабочий день.')
    doc.add_paragraph(
        '4.2. Работнику предоставляется ежегодный оплачиваемый отпуск продолжительностью 14 календарных дней.')
    doc.add_paragraph('4.3. Заявление на отпуск подается за 1 день до начала отпуска.')
    doc.add_paragraph('')

    doc.add_heading('5. СРОК ДЕЙСТВИЯ ДОГОВОРА', level=1)
    doc.add_paragraph('5.1. Настоящий договор вступает в силу с 01 января 2026 г. и действует до 31 декабря 2026 г.')
    doc.add_paragraph('')

    doc.add_paragraph('ПОДПИСИ СТОРОН:')
    doc.add_paragraph('Работодатель: _________________ /Иванов И.И./')
    doc.add_paragraph('Работник: _________________ /Петров П.П./')

    file_path = "data/contracts/test_employment_contract.docx"
    doc.save(file_path)
    print(f"✅ Тестовый договор создан: {file_path}")


if __name__ == "__main__":
    create_test_contract()