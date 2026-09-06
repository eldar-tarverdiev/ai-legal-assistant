import os
from typing import List
from pypdf import PdfReader
from docx import Document as DocxDocument
from langchain_core.documents import Document


def parse_pdf(file_path: str) -> List[Document]:
    """Парсит PDF и возвращает список документов."""
    reader = PdfReader(file_path)
    documents = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text.strip():
            documents.append(Document(
                page_content=text,
                metadata={
                    "source": os.path.basename(file_path),
                    "page": i + 1,
                    "type": "law"
                }
            ))

    return documents


def parse_docx(file_path: str) -> List[Document]:
    """Парсит DOCX и возвращает список документов."""
    doc = DocxDocument(file_path)
    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)

    text = "\n".join(full_text)

    return [Document(
        page_content=text,
        metadata={
            "source": os.path.basename(file_path),
            "type": "contract"
        }
    )]


def parse_document(file_path: str) -> List[Document]:
    """Универсальный парсер для разных форматов."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")