import os
from docx import Document


def create_sample_docx():
    # Создаем папку, если её нет
    os.makedirs("data/laws", exist_ok=True)

    doc = Document()
    doc.add_heading('Трудовой Кодекс (Выдержка)', 0)

    doc.add_heading('Статья 1. Общие положения', level=1)
    doc.add_paragraph('Каждый сотрудник имеет право на ежегодный оплачиваемый отпуск.')
    doc.add_paragraph('Продолжительность ежегодного основного оплачиваемого отпуска составляет 28 календарных дней.')

    doc.add_heading('Статья 2. Порядок предоставления отпуска', level=1)
    doc.add_paragraph(
        'Заявление на предоставление отпуска должно быть подано работодателю не позднее чем за 2 недели (14 календарных дней) до начала отпуска.')
    doc.add_paragraph('Работодатель обязан выплатить отпускные не позднее чем за 3 дня до начала отпуска.')

    file_path = "data/laws/sample_labor_code.docx"
    doc.save(file_path)
    print(f"✅ Тестовый документ успешно создан: {file_path}")


if __name__ == "__main__":
    create_sample_docx()